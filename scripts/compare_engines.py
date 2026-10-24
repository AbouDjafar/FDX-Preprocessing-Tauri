"""CI-only differential checks using generated data and the retained Python reference."""
import hashlib
import json
import posixpath
import subprocess
import sys
import tempfile
from collections import Counter
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from xml.etree import ElementTree
from zipfile import ZipFile

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'src'))
sys.path.insert(0, str(ROOT / 'tests'))
from conftest import project
from fdx_atelier.engine import analyze
from fdx_atelier.models import ChronologyRule, Decision, FieldRule, SegmentationRule, cell_key
from openpyxl import load_workbook


def exact_number(raw):
    value = Decimal(raw)
    assert value.is_finite(), f'Non-finite number: {raw!r}'
    return ('number', value)


def comparable(value, excel_type=None):
    if excel_type in ('f', 'e'):
        return ('formula' if excel_type == 'f' else 'error', value)
    if value is None or value == '':
        return ('empty', '')
    if isinstance(value, bool):
        return ('boolean', str(value).lower())
    if isinstance(value, (date, datetime)):
        if isinstance(value, datetime):
            assert not any((value.hour, value.minute, value.second, value.microsecond)), value
        return ('date', value.strftime('%Y-%m-%d'))
    if isinstance(value, (int, float, Decimal)):
        # Match the Python engine's Decimal(str(float)) conversion, not the
        # binary expansion Decimal(float). No rounding or tolerance is used.
        return exact_number(str(value))
    assert isinstance(value, str), f'Unexpected Python cell type: {type(value)}'
    return ('text', value)


def typed(value):
    kind = value['kind']
    raw = value['value']
    assert isinstance(kind, str) and isinstance(raw, str), value
    if kind in ('integer', 'decimal'):
        number = exact_number(raw)
        if kind == 'integer':
            assert number[1] == number[1].to_integral_value(), value
        # Both are numeric Excel storage. Text, booleans and dates stay distinct.
        return number
    if kind == 'empty':
        assert raw == '', value
        return ('empty', '')
    if kind == 'text' and raw == '':
        return ('empty', '')
    if kind == 'boolean':
        assert raw in ('true', 'false'), value
    elif kind == 'date':
        assert date.fromisoformat(raw).isoformat() == raw, value
    else:
        assert kind in ('text', 'formula', 'error'), value
    return (kind, raw)


def wire_value(value):
    kind, raw = comparable(value)
    if kind == 'number':
        kind = 'integer' if isinstance(value, int) else 'decimal'
        raw = str(raw)
    return {'kind': kind, 'value': raw}


def workbook_values(path):
    """Read exact XML numerics so even a float-roundtrip discrepancy fails."""
    ns = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    rel_ns = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
    numbers = {}
    with ZipFile(path) as archive:
        relations = ElementTree.fromstring(archive.read('xl/_rels/workbook.xml.rels'))
        targets = {node.attrib['Id']: node.attrib['Target'] for node in relations}
        workbook = ElementTree.fromstring(archive.read('xl/workbook.xml'))
        for sheet in workbook.findall('s:sheets/s:sheet', ns):
            target = targets[sheet.attrib[f'{{{rel_ns}}}id']]
            target = target.lstrip('/') if target.startswith('/') else posixpath.normpath(f'xl/{target}')
            xml = ElementTree.fromstring(archive.read(target))
            for cell in xml.findall('s:sheetData/s:row/s:c', ns):
                raw = cell.find('s:v', ns)
                if cell.attrib.get('t', 'n') == 'n' and cell.find('s:f', ns) is None and raw is not None and raw.text:
                    numbers[sheet.attrib['name'], cell.attrib['r']] = exact_number(raw.text)
    book = load_workbook(path, data_only=False)
    try:
        values = {}
        for sheet in book:
            for row in sheet:
                for cell in row:
                    value = comparable(cell.value, cell.data_type)
                    if value[0] == 'number':
                        value = numbers[sheet.title, cell.coordinate]
                    values[cell_key(sheet.title, cell.row, cell.column)] = value
        return values
    finally:
        book.close()


def semantic_changes(records, config, source, label):
    changes, noops, seen = {}, [], set()
    allowed = {m.column: m.field_code for m in config.mappings if m.field_code}
    if config.segmentation.enabled:
        seg = config.segmentation
        for column in (seg.name_column, seg.bp_column, seg.phone_column, seg.location_column):
            if column is not None:
                allowed.setdefault(column, f'COL_{column}')
    for record in records:
        key, row, column, code, before, after = record
        assert key not in seen, (label, 'duplicate change', key)
        seen.add(key)
        assert key == cell_key(config.sheet, row, column), (label, 'coordinates', record)
        assert row > config.header_row and column in allowed, (label, 'scope', record)
        assert code == allowed[column], (label, 'field code', record)
        assert key in source and before == source[key], (label, 'source mismatch', record, source.get(key))
        if before == after:
            noops.append(key)
        else:
            assert before[0] != 'formula', (label, 'formula modified', record)
            changes[key] = (row, column, code, before, after)
    return changes, noops


def check_comparator():
    """Guards for the oracle itself, executed by CI before invoking the CLI."""
    from types import SimpleNamespace

    config = SimpleNamespace(sheet='S', header_row=1, mappings=[SimpleNamespace(column=1, field_code='N')],
                             segmentation=SimpleNamespace(enabled=False))
    number = comparable(1250.12345678901)
    source = {'S!2:1': number}
    record = ('S!2:1', 2, 1, 'N', number, comparable(Decimal('1250.12345678901')))
    changes, noops = semantic_changes([record], config, source, 'selftest')
    assert not changes and noops == ['S!2:1']
    assert typed({'kind': 'decimal', 'value': '1250.12345678901'}) == number
    for after in (comparable('1250.12345678901'), comparable(True), comparable(date(2025, 1, 1)),
                  comparable(None), comparable(Decimal('1250.123456789010000000000001'))):
        changes, noops = semantic_changes([(*record[:-1], after)], config, source, 'selftest')
        assert len(changes) == 1 and not noops, after
    assert comparable('00123') != comparable(123)
    assert comparable('=1+1') != typed({'kind': 'formula', 'value': '=1+1'})
    assert typed({'kind': 'integer', 'value': '1'}) == typed({'kind': 'decimal', 'value': '1.000'})
    for records in ([record, record], [('S!2:1', 2, 1, 'N', comparable(0), comparable(0))]):
        try:
            semantic_changes(records, config, source, 'selftest')
        except AssertionError:
            pass
        else:
            raise AssertionError('Duplicate or false-before change was hidden as a noop')


def invoke(binary, *args):
    run = subprocess.run([str(binary), *map(str, args)], capture_output=True, text=True, encoding='utf-8')
    if run.returncode:
        raise AssertionError(f'CLI {args[0]} failed: {run.stderr}')
    return run.stdout


def configure_case(config, case):
    if case == 'decisions_chronology':
        config.rules['TEL'] = FieldRule(phone_policy='first')
        config.fields[0].max_length = 5
        config.rules['NOM'] = FieldRule(length_policy='truncate')
        config.chronology = [ChronologyRule(before='DATE', after='SIGN')]
        return {
            cell_key(config.sheet, 2, 1): Decision(action='set', value='  Exact manual value  '),
            cell_key(config.sheet, 2, 2): Decision(action='keep'),
            cell_key(config.sheet, 2, 4): Decision(action='set', value=datetime(2025, 9, 30)),
        }
    if case == 'cumulative_constraints':
        config.fields[0].max_length = 8
        config.fields[0].pattern = r'[A-Z]+'
        config.fields[0].allowed_values = ['Entreprise B']
        config.fields[2].minimum = Decimal('1500')
        config.fields[2].allowed_values = ['999']
        return {}
    if case == 'segmentation_conflicts':
        book = load_workbook(config.source_path)
        try:
            book[config.sheet]['A2'] = 'ETS ALPHA BP 123 Douala Tel 699123456'
            book[config.sheet]['A3'] = 'ETS BETA BP 456 Yaounde Tel +390498976845'
            book.save(config.source_path)
        finally:
            book.close()
        config.segmentation = SegmentationRule(enabled=True, source_column=1, name_column=1,
                                               bp_column=5, phone_column=2)
        return {cell_key(config.sheet, 2, 5): Decision(action='set', value='123')}
    assert case in ('report', 'first', 'all', 'clear'), case
    return {}


def check(binary, folder, policy, case=None):
    case = case or policy
    folder.mkdir()
    config = project.__wrapped__(folder)
    config.rules['TEL'] = FieldRule(phone_policy=policy)
    decisions = configure_case(config, case)
    baseline = hashlib.sha256(Path(config.source_path).read_bytes()).hexdigest()
    source = workbook_values(config.source_path)
    reference = analyze(config, decisions)
    config_file = folder / 'config.json'
    config_file.write_text(config.model_dump_json(), encoding='utf-8')
    decision_args = []
    if decisions:
        decision_file = folder / 'decisions.json'
        payload = {key: {'action': choice.action, 'value': wire_value(choice.value),
                         'allow_multiple': choice.allow_multiple, 'reason': choice.reason}
                   for key, choice in decisions.items()}
        decision_file.write_text(json.dumps(payload), encoding='utf-8')
        decision_args = ['--decisions', decision_file]
    result_file = folder / 'analysis.json'
    invoke(binary, 'analyze', '--config', config_file, *decision_args, '--output', result_file)
    actual = json.loads(result_file.read_text(encoding='utf-8'))
    for name in ('row_count', 'column_count', 'sheet', 'source_hash'):
        assert actual[name] == getattr(reference, name), (case, name)
    assert actual['source_hash'] == baseline
    expected_records = [(c.key, c.row, c.column, c.field_code, comparable(c.before), comparable(c.after))
                        for c in reference.changes]
    actual_records = [(c['key'], c['row'], c['column'], c['field_code'], typed(c['before']), typed(c['after']))
                      for c in actual['changes']]
    expected_changes, expected_noops = semantic_changes(expected_records, config, source, f'{case}: Python')
    actual_changes, actual_noops = semantic_changes(actual_records, config, source, f'{case}: Rust')
    assert actual_changes == expected_changes, (case, 'changes', expected_changes, actual_changes)
    final_values = {**source, **{key: change[-1] for key, change in expected_changes.items()}}
    # Counter keeps multiple constraints in the same cell; a set hid these.
    expected_findings = Counter((f.key, f.row, f.column, f.field_code, f.kind) for f in reference.findings)
    actual_findings = Counter((f['key'], f['row'], f['column'], f['field_code'], f['kind']) for f in actual['findings'])
    assert actual_findings == expected_findings, (case, 'findings', expected_findings, actual_findings)
    for finding in actual['findings']:
        key = finding['key']
        assert key == cell_key(config.sheet, finding['row'], finding['column']), (case, finding)
        assert typed(finding['original']) == source[key], (case, 'finding original', finding)
        assert typed(finding['current']) == final_values[key], (case, 'finding current', finding)
    if case == 'decisions_chronology':
        assert final_values[cell_key(config.sheet, 2, 1)] == comparable('  Exact manual value  ')
        assert final_values[cell_key(config.sheet, 2, 2)] == source[cell_key(config.sheet, 2, 2)]
        assert any(f.kind == 'chronology' and f.row == 2 for f in reference.findings)
    elif case == 'cumulative_constraints':
        assert sum(f.kind == 'constraint' and f.row == 2 and f.column == 3 for f in reference.findings) == 2
    elif case == 'segmentation_conflicts':
        assert final_values[cell_key(config.sheet, 2, 5)] == comparable('123')
        assert any(f.kind == 'segmentation_conflict' and f.row == 3 and f.column == 5 for f in reference.findings)
    target = folder / 'output.xlsx'
    invoke(binary, 'export', '--config', config_file, *decision_args, '--output', target, '--allow-issues')
    exported = workbook_values(target)
    for key, expected in final_values.items():
        assert exported.get(key, ('empty', '')) == expected, (case, 'exported value', key, expected, exported.get(key))
    book = load_workbook(target, data_only=False)
    try:
        assert book[config.sheet]['F3'].value == '=SUM(C2:C3)'
        assert book[config.sheet].freeze_panes == 'A2'
        assert book[config.sheet].auto_filter.ref == 'A1:G3'
        assert book['Conserver']['A2'].value == 12345.67
        for sheet in book:
            for row in sheet:
                for cell in row:
                    assert not cell.alignment.textRotation, (sheet.title, cell.coordinate)
                    if isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
                        assert '#,##' not in cell.number_format
        assert book['À vérifier'].max_row - 1 == len(actual['findings'])
    finally:
        book.close()
    assert hashlib.sha256(Path(config.source_path).read_bytes()).hexdigest() == baseline
    return {'case': case, 'policy': config.rules['TEL'].phone_policy, 'rows': actual['row_count'],
            'changes': len(actual_changes), 'findings': sum(actual_findings.values()),
            'raw_changes': {'python': len(reference.changes), 'rust': len(actual['changes'])},
            'semantic_noops': {'python': expected_noops, 'rust': actual_noops}}


if __name__ == '__main__':
    check_comparator()
    binary = Path(sys.argv[1]).resolve()
    with tempfile.TemporaryDirectory(prefix='fdx-oracle-') as temporary:
        results = [check(binary, Path(temporary) / policy, policy) for policy in ('report', 'first', 'all', 'clear')]
        results.extend(check(binary, Path(temporary) / case, 'report', case)
                       for case in ('decisions_chronology', 'cumulative_constraints', 'segmentation_conflicts'))
    destination = ROOT / 'qa' / 'parity.json'
    destination.parent.mkdir(exist_ok=True)
    destination.write_text(json.dumps({'ok': True, 'cases': results}, indent=2), encoding='utf-8')
    print(json.dumps({'ok': True, 'cases': len(results)}))
