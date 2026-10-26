"""Exercise the shipped Windows binary over its real Tauri bridge, never a mock."""
import hashlib
import json
import subprocess
import sys
import tempfile
import time
import urllib.error
import urllib.request
import uuid
from pathlib import Path

from openpyxl import load_workbook
from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.remote.client_config import ClientConfig
from selenium.webdriver.support.ui import Select, WebDriverWait
from selenium.webdriver.support import expected_conditions as EC


class TauriOptions(webdriver.ChromeOptions):
    def to_capabilities(self):
        return {'browserName': 'wry', 'tauri:options': {
            'application': str(Path(sys.argv[1]).resolve()), 'webviewOptions': {},
        }}


def call(driver, op, payload=None, session=None, fail=False):
    request = {'op': op, 'task_id': str(uuid.uuid4()), 'payload': payload or {}}
    if session:
        request.update(session_id=session['id'], revision=session['revision'])
    response = driver.execute_async_script('''
      const done = arguments[arguments.length - 1];
      window.__TAURI_INTERNALS__.invoke('fdx', { request: arguments[0], onEvent: null })
        .then(value => done({value})).catch(error => done({error}));
    ''', request)
    if fail:
        assert 'error' in response, (op, response)
        return response['error']
    assert 'error' not in response, (op, response)
    return response['value']


def check(driver, folder):
    wait = WebDriverWait(driver, 30)
    wait.until(lambda d: 'FDX' in d.find_element(By.TAG_NAME, 'body').text)
    assert driver.execute_script('return getComputedStyle(document.body).backgroundColor') != 'rgb(255, 255, 255)'
    driver.set_window_size(1366, 768)
    wait.until(EC.element_to_be_clickable((By.CSS_SELECTOR, '[data-testid="try-demo"]')))
    print('Native interface ready', flush=True)
    driver.save_screenshot(str(folder / 'native-start.png'))
    driver.find_element(By.CSS_SELECTOR, '[data-testid="try-demo"]').click()
    wait.until(EC.element_to_be_clickable((By.CSS_SELECTOR, '[data-testid="stage-association"]'))).click()
    confirm = '//button[normalize-space()="Confirmer toutes les associations"]'
    wait.until(EC.element_to_be_clickable((By.XPATH, confirm))).click()
    wait.until(EC.element_to_be_clickable((By.XPATH, '//dialog' + confirm))).click()
    wait.until(EC.invisibility_of_element_located((By.CSS_SELECTOR, 'dialog[open]')))
    wait.until(EC.element_to_be_clickable((By.XPATH, '//button[contains(., "Enregistrer et définir les règles")]'))).click()
    wait.until(EC.element_to_be_clickable((By.XPATH, '//button[contains(., "Confirmer et analyser toutes les lignes")]'))).click()
    wait.until(EC.visibility_of_element_located((By.CSS_SELECTOR, '[data-testid="finding"]')))
    driver.save_screenshot(str(folder / 'native-review.png'))
    print('Native analysis ready', flush=True)
    assert not driver.find_elements(By.CSS_SELECTOR, '[data-testid="global-error"]')
    finding = driver.find_elements(By.CSS_SELECTOR, '[data-testid="finding"]')[0]
    assert 'phone_multiple' in finding.text
    finding.find_element(By.XPATH, './/button[contains(., "Examiner et décider")]').click()
    decision = wait.until(EC.visibility_of_element_located((By.CSS_SELECTOR, 'dialog[open] select')))
    Select(decision).select_by_value('first')
    driver.find_element(By.XPATH, '//dialog//button[contains(., "Ajouter la décision au lot")]').click()
    wait.until(EC.element_to_be_clickable((By.XPATH, '//button[normalize-space()="Vérifier le lot"]'))).click()
    wait.until(EC.element_to_be_clickable((By.XPATH, '//button[normalize-space()="Analyser à nouveau"]')))
    driver.find_element(By.CSS_SELECTOR, '[data-testid="stage-export"]').click()
    wait.until(EC.visibility_of_element_located((By.XPATH, '//label[contains(., "J’autorise l’export avec les anomalies restantes")]/input'))).click()
    ui_output = folder / 'resultat interface.xlsx'
    # Substitute only the OS file picker; all application commands and Excel work remain native.
    driver.execute_script('''
      const original = window.__TAURI_INTERNALS__.invoke;
      const destination = arguments[0];
      window.__TAURI_INTERNALS__.invoke = (command, args, options) => {
        if (command === 'plugin:dialog|save') {
          window.__TAURI_INTERNALS__.invoke = original;
          return Promise.resolve(destination);
        }
        return original(command, args, options);
      };
    ''', str(ui_output))
    driver.find_element(By.CSS_SELECTOR, '[data-testid="export-button"]').click()
    wait.until(EC.element_to_be_clickable((By.XPATH, '//dialog//button[normalize-space()="Exporter et vérifier"]'))).click()
    wait.until(lambda _: ui_output.exists())
    ui_book = load_workbook(ui_output)
    assert ui_book['Marches fictifs']['B2'].value == '699123456'
    assert not any(row[6] == 'phone_multiple' for row in ui_book['À vérifier'].iter_rows(min_row=2, values_only=True))
    ui_book.close()
    wait.until(EC.element_to_be_clickable((By.XPATH, '//button[normalize-space()="Ouvrir le classeur"]')))
    driver.save_screenshot(str(folder / 'native-export.png'))
    print('Native UI export reopened and verified', flush=True)
    session = call(driver, 'demo')
    source = Path(session['config']['source_path'])
    baseline = hashlib.sha256(source.read_bytes()).hexdigest()
    config = session['config']
    for mapping in config['mappings']:
        mapping['confirmed'] = True
    session = call(driver, 'configure', {'config': config}, session)
    session = call(driver, 'analyze', session=session)
    assert session['summary']['row_count'] == 3
    findings = call(driver, 'results', {'view': 'findings', 'limit': 100}, session)['items']
    phone = next(f for f in findings if f['kind'] == 'phone_multiple')
    chosen = phone['choices'][0]
    session = call(driver, 'decide', {'decisions': {phone['key']: {'action': 'set', 'value': chosen, 'allow_multiple': False, 'reason': 'Recette native'}}}, session)
    output = folder / 'resultat natif.xlsx'
    call(driver, 'export', {'path': str(output), 'allow_issues': True, 'overwrite': False}, session, fail=True)
    assert not output.exists()
    session = call(driver, 'undo', session=session)
    session = call(driver, 'redo', session=session)
    session = call(driver, 'analyze', session=session)
    result = call(driver, 'export', {'path': str(output), 'allow_issues': True, 'overwrite': False}, session)
    assert Path(result['path']).exists()
    workbook = load_workbook(output)
    assert str(workbook['Marches fictifs']['B2'].value) == chosen['value']
    assert workbook['Marches fictifs']['E2'].value == '000123'
    assert workbook['À vérifier'].max_row - 1 == session['summary']['findings']
    for sheet in workbook:
        for row in sheet:
            for cell in row:
                assert not cell.alignment.textRotation
    workbook.close()
    assert hashlib.sha256(source.read_bytes()).hexdigest() == baseline
    project = folder / 'projet.fdxproj'
    call(driver, 'project_save', {'path': str(project)}, session)
    restored = call(driver, 'project_open', {'path': str(project)})
    assert restored['dirty'] is True
    restored = call(driver, 'analyze', session=restored)
    assert restored['summary']['findings'] == session['summary']['findings']
    print('Native project reopened and verified', flush=True)
    return {'ok': True, 'rows': 3, 'native_ui': True, 'native_ui_export': True, 'native_bridge': True, 'source_unchanged': True, 'export_reopened': True, 'project_reopened': True}


if __name__ == '__main__':
    destination = Path('qa/native')
    destination.mkdir(parents=True, exist_ok=True)
    log = (destination / 'driver.log').open('w', encoding='utf-8')
    server = subprocess.Popen(['tauri-driver', '--native-driver', str(Path(sys.argv[2]).resolve())], stdout=log, stderr=subprocess.STDOUT)
    driver = None
    try:
        for attempt in range(30):
            if server.poll() is not None:
                raise RuntimeError('tauri-driver exited before becoming ready; inspect driver.log')
            try:
                with urllib.request.urlopen('http://127.0.0.1:4444/status', timeout=2) as response:
                    if json.load(response).get('value', {}).get('ready'):
                        break
            except (OSError, urllib.error.URLError, json.JSONDecodeError):
                pass
            time.sleep(1)
        else:
            raise RuntimeError('WebDriver did not become ready; inspect driver.log')
        print('WebDriver ready; creating one native application session', flush=True)
        driver = webdriver.Remote(
            command_executor='http://127.0.0.1:4444', options=TauriOptions(),
            client_config=ClientConfig('http://127.0.0.1:4444', timeout=60),
        )
        print('Native application session created', flush=True)
        driver.set_script_timeout(45)
        report = check(driver, destination.resolve())
        (destination / 'verification.json').write_text(json.dumps(report, indent=2), encoding='utf-8')
        print(json.dumps(report))
    except BaseException:
        if driver:
            try:
                driver.save_screenshot(str(destination / 'native-failure.png'))
                (destination / 'failure.html').write_text(driver.page_source[:200000], encoding='utf-8')
            except Exception:
                pass
        raise
    finally:
        if driver:
            try:
                driver.quit()
            except Exception:
                pass
        server.terminate()
        try:
            server.wait(timeout=10)
        except subprocess.TimeoutExpired:
            server.kill()
        log.close()
