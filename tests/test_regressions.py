import threading
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from openpyxl import load_workbook
from openpyxl.formatting.rule import Rule
from openpyxl.styles import Alignment
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.styles.numbers import NumberFormat

from fdx_atelier.dictionary import suggest_mappings
from fdx_atelier.engine import analyze
from fdx_atelier.exporter import export_excel, numeric_format
from fdx_atelier.models import Cancelled, Decision, FieldSpec, Role, SegmentationRule, cell_key
from fdx_atelier.workbooks import WorkbookError, inspect_workbook, load_preserving


def test_matching_reserves_exact_phone_and_identifier():
    fields = [
        FieldSpec(code="TEL", label="Téléphone de l'entreprise", role=Role.PHONE),
        FieldSpec(code="NIU", label="NIU", role=Role.IDENTIFIER),
        FieldSpec(code="NOM", label="Nom de l'entreprise", role=Role.COMPANY),
    ]
    mappings = suggest_mappings(
        ["Catégorie", "Nom de l'entreprise", "NIU", "Téléphone de l'entreprise"], fields
    )
    assert [m.field_code for m in mappings] == [None, "NOM", "NIU", "TEL"]
    assert not any(m.confirmed for m in mappings)


def test_workbook_import_can_be_cancelled(project):
    cancel = threading.Event()
    cancel.set()
    with pytest.raises(Cancelled):
        inspect_workbook(project.source_path, cancel=cancel)


def test_mislabeled_numeric_date_is_not_lost(project, tmp_path):
    wb = load_workbook(project.source_path)
    wb.active["F2"] = 9999999999
    wb.active["F2"].number_format = "dd/mm/yyyy"
    wb.active["F3"] = ""
    wb.save(project.source_path)
    wb.close()
    wb = load_preserving(project.source_path)
    assert wb.active["F2"].value == 9999999999
    wb.close()
    result = analyze(project)
    target = export_excel(project, result, tmp_path / "valeurs.xlsx", allow_issues=True)
    wb = load_workbook(target)
    assert wb.active["F2"].value == 9999999999
    assert wb.active["F3"].value in {None, ""}
    wb.close()


def test_inherited_and_conditional_formats_are_harmonized(project, tmp_path):
    wb = load_workbook(project.source_path)
    sheet = wb.active
    sheet.column_dimensions["C"].alignment = Alignment(textRotation=90)
    sheet.row_dimensions[3].alignment = Alignment(textRotation=255)
    sheet.column_dimensions["C"].number_format = "#,##0.00"
    sheet.conditional_formatting.add(
        "C2:C3",
        Rule(
            type="cellIs",
            operator="greaterThan",
            formula=["0"],
            dxf=DifferentialStyle(numFmt=NumberFormat(numFmtId=164, formatCode="#,##0.00")),
        ),
    )
    wb.save(project.source_path)
    wb.close()
    result = analyze(project)
    target = export_excel(project, result, tmp_path / "formats.xlsx", allow_issues=True)
    wb = load_workbook(target)
    sheet = wb.active
    assert sheet.column_dimensions["C"].alignment.textRotation == 0
    assert sheet.row_dimensions[3].alignment.textRotation == 0
    assert "," not in sheet.column_dimensions["C"].number_format
    for conditional in sheet.conditional_formatting:
        assert "," not in sheet.conditional_formatting[conditional][0].dxf.numFmt.formatCode
    wb.close()


def test_report_collisions_are_preserved_by_default(project, tmp_path):
    wb = load_workbook(project.source_path)
    wb.create_sheet("À vérifier")["A1"] = "Rapport existant"
    wb.save(project.source_path)
    wb.close()
    result = analyze(project)
    target = export_excel(project, result, tmp_path / "rapports.xlsx", allow_issues=True)
    wb = load_workbook(target)
    assert wb["À vérifier"]["A1"].value == "Rapport existant"
    assert wb["À vérifier (2)"].max_row - 1 == len(result.findings)
    wb.close()


def test_numeric_display_preserves_small_decimals_and_percent():
    assert numeric_format(1.234e-20, "General").count("#") == 23
    assert "E+00" in numeric_format(1e-50, "General")
    assert numeric_format(0.125, "0.0%") == "0.#%"
    assert numeric_format(12, '0"%"') == "0"


def test_segmentation_preserves_conflicting_destinations_until_decided(project, tmp_path):
    wb = load_workbook(project.source_path)
    wb.active["A2"] = "ETS ALPHA BP: 123 Douala Tel: 699123456"
    wb.save(project.source_path)
    wb.close()

    project.segmentation = SegmentationRule(
        enabled=True, source_column=1, name_column=1, bp_column=5, phone_column=2, location_column=6
    )
    result = analyze(project)
    bp_key = cell_key(project.sheet, 2, 5)
    conflict = next(f for f in result.findings if f.key == bp_key and f.kind == "segmentation_conflict")
    assert conflict.proposed == "123"
    assert not any(c.key == bp_key for c in result.changes)
    fixed = analyze(project, {bp_key: Decision(action="set", value="123")})
    assert not any(f.key == bp_key and f.kind == "segmentation_conflict" for f in fixed.findings)
    assert next(c for c in fixed.changes if c.key == bp_key).after == "123"
    target = export_excel(project, fixed, tmp_path / "segmentation.xlsx", allow_issues=True)
    wb = load_workbook(target)
    assert wb.active["A2"].value == "ETS ALPHA"
    assert wb.active["E2"].value == "123"
    assert any(
        row[3] == "ETS ALPHA BP: 123 Douala Tel: 699123456"
        for row in wb["Nom_original"].iter_rows(min_row=2, values_only=True)
    )
    wb.close()


@pytest.mark.parametrize("payload", ["dtd", "extension"])
def test_unsafe_xml_and_prefixed_extensions_are_rejected(project, payload):
    with ZipFile(project.source_path) as archive:
        entries = {name: archive.read(name) for name in archive.namelist()}
    name = "xl/worksheets/sheet1.xml"
    if payload == "dtd":
        entries[name] = entries[name].replace(
            b"<worksheet", b'<!DOCTYPE worksheet [<!ENTITY x "interdit">]><worksheet', 1
        )
    else:
        entries[name] = entries[name].replace(
            b"</worksheet>",
            b'<s:extLst xmlns:s="http://schemas.openxmlformats.org/spreadsheetml/2006/main"/></worksheet>',
        )
    with ZipFile(project.source_path, "w", compression=ZIP_DEFLATED) as archive:
        for name, content in entries.items():
            archive.writestr(name, content)
    with pytest.raises(WorkbookError):
        inspect_workbook(project.source_path)
