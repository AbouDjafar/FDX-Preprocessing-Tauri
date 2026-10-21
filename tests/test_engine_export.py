import threading
from datetime import datetime
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from fdx_atelier.engine import analyze
from fdx_atelier.exporter import export_excel, remove_grouping
from fdx_atelier.models import Cancelled, ChronologyRule, Decision, FieldRule, cell_key
from fdx_atelier.workbooks import sha256


def test_full_analysis_and_decision_precedence(project):
    digest = sha256(project.source_path)
    result = analyze(project)
    assert result.row_count == 2
    assert any(f.kind == "phone_multiple" for f in result.findings)
    assert any(f.kind == "date_invalid" for f in result.findings)
    assert any(c.after == Decimal("1234.50") for c in result.changes)
    project.rules["TEL"] = FieldRule(phone_policy="first")
    key = cell_key(project.sheet, 2, 2)
    first = analyze(project)
    assert next(c for c in first.changes if c.key == key).after == "+237699123456"
    overridden = analyze(project, {key: Decision(action="keep")})
    assert not any(c.key == key for c in overridden.changes)
    assert any(f.kind == "phone_multiple" for f in overridden.findings)
    assert sha256(project.source_path) == digest


def test_accept_date_and_multiple_phones(project):
    decisions = {
        cell_key(project.sheet, 2, 4): Decision(action="set", value=datetime(2025, 9, 30)),
        cell_key(project.sheet, 2, 2): Decision(
            action="set", value="+237699123456/677123456", allow_multiple=True
        ),
    }
    result = analyze(project, decisions)
    assert not any(f.kind in {"phone_multiple", "date_invalid"} for f in result.findings)


def test_pandera_constraints_and_chronology(project):
    project.fields[0].allowed_values = ["Entreprise B"]
    project.fields[2].minimum = Decimal(1500)
    project.chronology = [ChronologyRule(before="DATE", after="SIGN")]
    result = analyze(project, {cell_key(project.sheet, 2, 4): Decision(action="set", value="2025-09-30")})
    assert any(f.kind == "constraint" and f.field_code == "NOM" for f in result.findings)
    assert any(f.kind == "constraint" and f.field_code == "MON" for f in result.findings)
    assert any(f.kind == "chronology" for f in result.findings)


def test_export_conserves_unselected_values_formulas_and_precision(project, tmp_path):
    before = sha256(project.source_path)
    result = analyze(project)
    target = tmp_path / "résultat.xlsx"
    export_excel(project, result, target, allow_issues=True)
    workbook = load_workbook(target, data_only=False)
    assert workbook.sheetnames == ["Données", "Conserver", "Contrôle", "Nom_original", "À vérifier"]
    assert workbook["Données"]["F3"].value == "=SUM(C2:C3)"
    assert workbook["Données"]["E2"].value == "00123"
    assert workbook["Données"]["C3"].value == 1250.12345678901
    assert workbook["Données"]["C2"].value == 1234.5
    assert workbook["Données"].freeze_panes == "A2"
    assert workbook["À vérifier"].max_row - 1 == len(result.findings)
    for sheet in workbook:
        for row in sheet:
            for cell in row:
                assert cell.alignment.textRotation == 0
                if isinstance(cell.value, (int, float)):
                    assert "," not in cell.number_format
    workbook.close()
    assert sha256(project.source_path) == before


def test_plain_output_and_stale_result(project, tmp_path):
    project.harmonized = False
    result = analyze(project)
    export_excel(project, result, tmp_path / "simple.xlsx", allow_issues=True)
    workbook = load_workbook(tmp_path / "simple.xlsx")
    assert workbook.sheetnames == ["Données", "Conserver"]
    workbook.close()
    project.rules["TEL"] = FieldRule(phone_policy="first")
    with pytest.raises(ValueError, match="obsolète"):
        export_excel(project, result, tmp_path / "obsolete.xlsx", allow_issues=True)


def test_export_errors_are_atomic(project, tmp_path, monkeypatch):
    result = analyze(project)
    with pytest.raises(ValueError, match="source"):
        export_excel(project, result, project.source_path, allow_issues=True, overwrite=True)
    with pytest.raises(ValueError, match="anomalies"):
        export_excel(project, result, tmp_path / "refus.xlsx")
    digest = sha256(project.source_path)
    import fdx_atelier.exporter as module

    def locked(*args):
        raise PermissionError("Fichier ouvert dans Excel")

    monkeypatch.setattr(module.os, "replace", locked)
    with pytest.raises(PermissionError):
        export_excel(project, result, tmp_path / "verrouille.xlsx", allow_issues=True)
    assert not list(tmp_path.glob(".fdx-*.xlsx"))
    assert not (tmp_path / "verrouille.xlsx").exists()
    assert sha256(project.source_path) == digest


def test_source_change_and_cancel(project, tmp_path):
    result = analyze(project)
    workbook = load_workbook(project.source_path)
    workbook.active["F2"] = "modification externe"
    workbook.save(project.source_path)
    workbook.close()
    with pytest.raises(ValueError, match="obsolète"):
        export_excel(project, result, tmp_path / "stale.xlsx", allow_issues=True)
    event = threading.Event()
    event.set()
    with pytest.raises(Cancelled):
        analyze(project, cancel=event)


def test_formula_like_user_text_is_not_executed(project, tmp_path):
    result = analyze(project, {cell_key(project.sheet, 2, 1): Decision(action="set", value="=1+1")})
    target = export_excel(project, result, tmp_path / "literal.xlsx", allow_issues=True)
    workbook = load_workbook(target)
    assert workbook.active["A2"].value == "=1+1"
    assert workbook.active["A2"].data_type == "s"
    workbook.close()


def test_duplicate_mapping_and_grouping_format(project):
    project.mappings[1].field_code = "NOM"
    with pytest.raises(ValueError, match="deux colonnes"):
        analyze(project)
    assert remove_grouping("#,##0.00;[Red]-#,##0.00") == "###0.00;[Red]-###0.00"
