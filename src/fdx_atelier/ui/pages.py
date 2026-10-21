from __future__ import annotations

from decimal import Decimal
from pathlib import Path

from openpyxl.utils import get_column_letter
from PySide6.QtCore import QLocale, QSignalBlocker, Qt, QUrl
from PySide6.QtGui import QDesktopServices
from PySide6.QtWidgets import (
    QAbstractItemView,
    QCheckBox,
    QComboBox,
    QFileDialog,
    QFormLayout,
    QGroupBox,
    QHBoxLayout,
    QHeaderView,
    QLabel,
    QLineEdit,
    QPlainTextEdit,
    QPushButton,
    QSpinBox,
    QTableWidget,
    QTableWidgetItem,
    QToolButton,
    QVBoxLayout,
    QWidget,
)

from ..dictionary import suggest_mappings
from ..models import (
    ChronologyRule,
    ColumnMapping,
    DataType,
    FieldRule,
    FieldSpec,
    Role,
    SegmentationRule,
    display,
)
from .icons import icon

TYPE_LABELS = {
    DataType.TEXT: "Texte",
    DataType.INTEGER: "Entier",
    DataType.DECIMAL: "Décimal",
    DataType.DATE: "Date",
    DataType.BOOLEAN: "Booléen",
}
ROLE_LABELS = {
    Role.STANDARD: "Standard",
    Role.PHONE: "Téléphone",
    Role.IDENTIFIER: "Identifiant",
    Role.COMPANY: "Entreprise",
}


def combo(options, selected=None):
    box = QComboBox()
    for value, label in options.items():
        box.addItem(label, value)
    idx = box.findData(selected)
    if idx >= 0:
        box.setCurrentIndex(idx)
    return box


def item(text, editable=True):
    result = QTableWidgetItem(str(text) if text is not None else "")
    if not editable:
        result.setFlags(result.flags() & ~Qt.ItemFlag.ItemIsEditable)
    return result


def button(text, callback, name=""):
    result = QPushButton(text)
    result.setObjectName(name)
    result.setAccessibleName(text)
    result.clicked.connect(callback)
    return result


def icon_button(label, icon, callback):
    result = QToolButton()
    result.setIcon(result.style().standardIcon(icon))
    result.setToolTip(label)
    result.setAccessibleName(label)
    result.setFixedSize(38, 38)
    result.clicked.connect(callback)
    return result


def setup_table(table):
    table.setAlternatingRowColors(True)
    table.setSelectionBehavior(QAbstractItemView.SelectionBehavior.SelectRows)
    table.verticalHeader().setVisible(False)
    table.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeMode.Interactive)
    table.horizontalHeader().setStretchLastSection(True)
    table.setWordWrap(False)
    table.verticalHeader().setDefaultSectionSize(42)


class ImportPage(QWidget):
    def __init__(self, app):
        super().__init__()
        self.app = app
        self.setAcceptDrops(True)
        layout = QVBoxLayout(self)
        group = QGroupBox("Quels fichiers souhaitez-vous préparer ?")
        form = QFormLayout(group)
        self.source = QLineEdit()
        self.source.setReadOnly(True)
        self.source.setObjectName("sourcePath")
        self.source.setPlaceholderText("Aucun fichier Excel sélectionné")
        self.dictionary = QLineEdit()
        self.dictionary.setReadOnly(True)
        self.dictionary.setPlaceholderText("Choisir les règles attendues (.md ou .txt)")
        self.dictionary.setToolTip(
            "Le dictionnaire définit les colonnes, leurs types et les valeurs attendues."
        )
        form.addRow("Classeur Excel", self.file_row(self.source, self.browse_source))
        form.addRow("Dictionnaire", self.file_row(self.dictionary, self.browse_dictionary))
        self.sheet = QComboBox()
        self.sheet.setObjectName("sheetSelector")
        self.header = QSpinBox()
        self.header.setRange(1, 1000)
        self.header.setKeyboardTracking(False)
        self.header.setObjectName("headerRow")
        self.header.setToolTip("Numéro de la ligne contenant les noms des colonnes dans Excel.")
        form.addRow("Feuille de données", self.sheet)
        form.addRow("Ligne d'en-tête", self.header)
        layout.addWidget(group)
        demo = button("Essayer un exemple", app.load_demo, "demoButton")
        demo.setIcon(icon("flask-conical"))
        demo.setToolTip("Ouvrir un petit classeur fictif, sans utiliser vos fichiers.")
        layout.addWidget(demo, 0, Qt.AlignmentFlag.AlignLeft)
        self.summary = QLabel("Aucun classeur sélectionné")
        self.summary.setObjectName("sourceSummary")
        layout.addWidget(self.summary)
        self.preview = QTableWidget()
        self.preview.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
        setup_table(self.preview)
        self.preview.setMinimumHeight(180)
        layout.addWidget(self.preview, 1)
        self.sheet.currentIndexChanged.connect(self.show_sheet)
        self.header.valueChanged.connect(self.header_changed)

    @staticmethod
    def file_row(line, callback):
        widget = QWidget()
        row = QHBoxLayout(widget)
        row.setContentsMargins(0, 0, 0, 0)
        row.addWidget(line, 1)
        browse = button("Choisir…", callback)
        browse.setIcon(icon("folder-open"))
        row.addWidget(browse)
        return widget

    def browse_source(self):
        path, _ = QFileDialog.getOpenFileName(self, "Ouvrir le classeur", "", "Classeur Excel (*.xlsx)")
        if path:
            self.app.set_source(path)

    def browse_dictionary(self):
        path, _ = QFileDialog.getOpenFileName(self, "Ouvrir le dictionnaire", "", "Dictionnaire (*.md *.txt)")
        if path:
            self.app.set_dictionary(path)

    def dragEnterEvent(self, event):
        if event.mimeData().hasUrls() and not self.app.busy:
            event.acceptProposedAction()

    def dropEvent(self, event):
        paths = [Path(url.toLocalFile()) for url in event.mimeData().urls()]
        excels = [p for p in paths if p.suffix.lower() == ".xlsx"]
        dictionaries = [p for p in paths if p.suffix.lower() in {".md", ".txt"}]
        if len(excels) > 1 or len(dictionaries) > 1:
            self.app.show_error("Choisissez un seul classeur et un seul dictionnaire à la fois.")
            return
        if excels:
            self.app.set_source(str(excels[0]))
        if dictionaries:
            self.app.set_dictionary(str(dictionaries[0]))
        event.acceptProposedAction()

    def load(self):
        path = self.app.config.original_path or self.app.config.source_path
        self.source.setText(Path(path).name if path else "")
        self.dictionary.setText(
            Path(self.app.config.dictionary_path).name if self.app.config.dictionary_path else ""
        )
        self.source.setToolTip(path)
        self.dictionary.setToolTip(
            self.app.config.dictionary_path or "Le dictionnaire définit les colonnes et leurs contraintes."
        )
        self.preview.setVisible(bool(self.app.metadata))
        self.sheet.setEnabled(bool(self.app.metadata))
        self.header.setEnabled(bool(self.app.metadata))
        with QSignalBlocker(self.header):
            self.header.setValue(self.app.config.header_row)
        with QSignalBlocker(self.sheet):
            self.sheet.clear()
            for sheet in (self.app.metadata or {}).get("sheets", []):
                self.sheet.addItem(sheet["name"])
            self.sheet.setCurrentText(self.app.config.sheet)
        self.show_sheet()

    def header_changed(self):
        if self.app.config.source_path and not self.app.busy:
            self.app.config.header_row = self.header.value()
            self.app.config.mappings = []
            self.app.invalidate()
            self.app.refresh_source()

    def show_sheet(self):
        name = self.sheet.currentText()
        sheets = (self.app.metadata or {}).get("sheets", [])
        sheet = next((s for s in sheets if s["name"] == name), None)
        if sheet is None:
            self.preview.clear()
            self.preview.setRowCount(0)
            self.summary.setText("Aucun fichier sélectionné")
            return
        if name != self.app.config.sheet:
            self.app.config.sheet = name
            self.app.config.mappings = []
            self.app.invalidate()
        self.summary.setText(
            f"{sheet['rows'] - self.header.value():,} lignes de données    ·    "
            f"{sheet['columns']} colonnes".replace(",", " ")
        )
        self.preview.clear()
        self.preview.setColumnCount(sheet["columns"])
        self.preview.setHorizontalHeaderLabels(
            [f"{get_column_letter(i)} · {display(h)}" for i, h in enumerate(sheet["headers"], 1)]
        )
        self.preview.setRowCount(len(sheet["preview"]))
        for row, values in enumerate(sheet["preview"]):
            for col, value in enumerate(values):
                self.preview.setItem(row, col, item(display(value), False))
        for col in range(sheet["columns"]):
            self.preview.setColumnWidth(col, 180)

    def collect(self):
        self.app.config.header_row = self.header.value()
        self.app.config.sheet = self.sheet.currentText()


class DictionaryPage(QWidget):
    HEADERS = [
        "Code",
        "Libellé",
        "Stockage",
        "Rôle",
        "Taille max",
        "Format date",
        "Obligatoire",
        "Motif",
        "Valeurs autorisées (une par ligne)",
        "Minimum",
        "Maximum",
    ]

    def __init__(self, app):
        super().__init__()
        self.app = app
        layout = QVBoxLayout(self)
        tools = QHBoxLayout()
        tools.addWidget(button("Ajouter un champ", self.add))
        tools.addWidget(button("Supprimer la sélection", self.remove))
        tools.addStretch()
        layout.addLayout(tools)
        self.table = QTableWidget(0, len(self.HEADERS))
        self.table.setObjectName("dictionaryTable")
        self.table.setHorizontalHeaderLabels(self.HEADERS)
        setup_table(self.table)
        self.table.setColumnWidth(0, 165)
        self.table.setColumnWidth(1, 280)
        layout.addWidget(self.table, 1)
        self.notes = QPlainTextEdit()
        self.notes.setReadOnly(True)
        self.notes.setMaximumHeight(100)
        layout.addWidget(self.notes)

    def add(self, spec=None):
        if not isinstance(spec, FieldSpec):
            spec = FieldSpec(code=f"CHAMP_{self.table.rowCount() + 1}", label="Nouveau champ")
        row = self.table.rowCount()
        self.table.insertRow(row)
        values = [
            spec.code,
            spec.label,
            "",
            "",
            spec.max_length,
            spec.date_format,
            "",
            spec.pattern,
            "\n".join(spec.allowed_values),
            spec.minimum,
            spec.maximum,
        ]
        for col, value in enumerate(values):
            self.table.setItem(row, col, item(value))
        self.table.setCellWidget(row, 2, combo(TYPE_LABELS, spec.dtype))
        self.table.setCellWidget(row, 3, combo(ROLE_LABELS, spec.role))
        required = QCheckBox()
        required.setChecked(spec.required)
        self.table.setCellWidget(row, 6, required)
        allowed = QPlainTextEdit("\n".join(spec.allowed_values))
        allowed.setMaximumHeight(70)
        allowed.setAccessibleName("Valeurs autorisées, une par ligne")
        self.table.setCellWidget(row, 8, allowed)

    def remove(self):
        for row in sorted({index.row() for index in self.table.selectedIndexes()}, reverse=True):
            self.table.removeRow(row)

    def load(self):
        self.table.setRowCount(0)
        for spec in self.app.config.fields:
            self.add(spec)
        self.notes.setPlainText("\n".join(self.app.dictionary_notes))
        self.notes.setVisible(bool(self.app.dictionary_notes))

    def collect(self):
        fields = []
        for row in range(self.table.rowCount()):

            def text(col):
                return self.table.item(row, col).text().strip()

            allowed = [
                v.strip() for v in self.table.cellWidget(row, 8).toPlainText().splitlines() if v.strip()
            ]
            fields.append(
                FieldSpec(
                    code=text(0),
                    label=text(1),
                    dtype=self.table.cellWidget(row, 2).currentData(),
                    role=self.table.cellWidget(row, 3).currentData(),
                    max_length=int(text(4)) if text(4) else None,
                    date_format=text(5),
                    required=self.table.cellWidget(row, 6).isChecked(),
                    pattern=text(7),
                    allowed_values=allowed,
                    minimum=Decimal(text(9).replace(",", ".")) if text(9) else None,
                    maximum=Decimal(text(10).replace(",", ".")) if text(10) else None,
                )
            )
        if len({f.code for f in fields}) != len(fields) or any(not f.code for f in fields):
            raise ValueError("Chaque champ doit avoir un code unique non vide.")
        self.app.config.fields = fields


class MappingPage(QWidget):
    def __init__(self, app):
        super().__init__()
        self.app = app
        layout = QVBoxLayout(self)
        tools = QHBoxLayout()
        self.recalculate = button("Recalculer les propositions", lambda: self.load(force=True))
        self.recalculate.setIcon(icon("undo-2"))
        tools.addWidget(self.recalculate)
        tools.addStretch()
        layout.addLayout(tools)
        self.table = QTableWidget(0, 7)
        self.table.setObjectName("mappingTable")
        self.table.setHorizontalHeaderLabels(
            [
                "Validée",
                "Colonne",
                "Colonne Excel",
                "Exemple de valeur",
                "Champ attendu",
                "Correspondance",
                "Format",
            ]
        )
        setup_table(self.table)
        self.table.setColumnWidth(2, 205)
        self.table.setColumnWidth(3, 175)
        self.table.setColumnWidth(4, 245)
        layout.addWidget(self.table, 1)

    def load(self, force=False):
        sheet = self.app.sheet_info()
        if not sheet:
            return
        mappings = self.app.config.mappings
        if force or not mappings:
            mappings = suggest_mappings(sheet["headers"], self.app.config.fields)
            self.app.config.mappings = mappings
        else:
            by_column = {m.column: m for m in mappings}
            mappings = [
                by_column.get(col, ColumnMapping(column=col)) for col in range(1, sheet["columns"] + 1)
            ]
            self.app.config.mappings = mappings
        self.table.clearContents()
        self.table.setRowCount(len(mappings))
        for row, mapping in enumerate(mappings):
            valid = QCheckBox()
            valid.setChecked(mapping.confirmed)
            self.table.setCellWidget(row, 0, valid)
            self.table.setItem(row, 1, item(f"{get_column_letter(mapping.column)} ({mapping.column})", False))
            self.table.setItem(row, 2, item(display(sheet["headers"][mapping.column - 1]), False))
            sample = next(
                (
                    display(values[mapping.column - 1])
                    for values in sheet["preview"]
                    if mapping.column <= len(values) and values[mapping.column - 1] is not None
                ),
                "",
            )
            self.table.setItem(row, 3, item(sample, False))
            box = combo(
                {None: "Ignorer", **{f.code: f"{f.code} · {f.label}" for f in self.app.config.fields}},
                mapping.field_code,
            )
            box.currentIndexChanged.connect(lambda _, check=valid: check.setChecked(False))
            self.table.setCellWidget(row, 4, box)
            self.table.setItem(
                row, 5, item(f"{mapping.score:.0f} %" if mapping.field_code else "À choisir", False)
            )
            spec = next((f for f in self.app.config.fields if f.code == mapping.field_code), None)
            formats = combo(TYPE_LABELS, spec.dtype if spec else DataType.TEXT)
            formats.setEnabled(spec is not None)
            formats.setAccessibleName(f"Format de la colonne {mapping.column}")
            self.table.setCellWidget(row, 6, formats)
            box.currentIndexChanged.connect(lambda _, r=row: self.match_changed(r))
        self.set_advanced(self.app.advanced)

    def match_changed(self, row):
        code = self.table.cellWidget(row, 4).currentData()
        spec = next((f for f in self.app.config.fields if f.code == code), None)
        box = self.table.cellWidget(row, 6)
        box.setEnabled(spec is not None)
        if spec:
            box.setCurrentIndex(box.findData(spec.dtype))

    def set_advanced(self, advanced):
        self.recalculate.setVisible(advanced)
        for col in (0, 1, 5):
            self.table.setColumnHidden(col, not advanced)
        for row in range(self.table.rowCount()):
            self.table.cellWidget(row, 0).setVisible(advanced)
            fields = self.table.cellWidget(row, 4)
            for index in range(1, fields.count()):
                spec = next((f for f in self.app.config.fields if f.code == fields.itemData(index)), None)
                if spec:
                    fields.setItemText(index, f"{spec.code} · {spec.label}" if advanced else spec.label)

    def confirm(self):
        for row in range(self.table.rowCount()):
            self.table.cellWidget(row, 0).setChecked(bool(self.table.cellWidget(row, 4).currentData()))
        self.app.statusBar().showMessage(
            "Associations validées ; les doublons seront contrôlés avant l'analyse."
        )

    def collect(self):
        old = self.app.config.mappings
        self.app.config.mappings = [
            ColumnMapping(
                column=old[row].column,
                field_code=self.table.cellWidget(row, 4).currentData(),
                confirmed=self.table.cellWidget(row, 0).isChecked(),
                score=old[row].score,
            )
            for row in range(self.table.rowCount())
        ]
        for row, mapping in enumerate(self.app.config.mappings):
            spec = next((f for f in self.app.config.fields if f.code == mapping.field_code), None)
            if spec:
                spec.dtype = DataType(self.table.cellWidget(row, 6).currentData())


class RulesPage(QWidget):
    def __init__(self, app):
        super().__init__()
        self.app = app
        self.current_code = None
        layout = QVBoxLayout(self)
        label = QLabel("Quelle colonne souhaitez-vous régler ?")
        label.setProperty("heading", True)
        layout.addWidget(label)
        self.field = QComboBox()
        self.field.setObjectName("ruleField")
        self.field.setToolTip("Les réglages concernent uniquement la colonne sélectionnée.")
        layout.addWidget(self.field)
        self.general = QGroupBox("Traitement du champ")
        form = QFormLayout(self.general)
        self.normalize = QCheckBox("Normaliser les espaces")
        self.convert = QCheckBox("Convertir vers le type validé")
        self.convert.setToolTip(
            "Préparer les conversions compatibles avec le format confirmé, sans corriger les cas ambigus."
        )
        form.addRow(self.normalize)
        form.addRow(self.convert)
        self.length = combo(
            {
                "report": "Conserver et signaler",
                "truncate": "Tronquer à la taille maximale",
                "clear": "Vider la cellule",
            }
        )
        self.invalid = combo(
            {"report": "Conserver et signaler", "clear": "Vider si le motif n'est pas respecté"}
        )
        form.addRow("Texte trop long", self.length)
        form.addRow("Motif non conforme", self.invalid)
        self.replacements = QTableWidget(0, 2)
        self.replacements.setHorizontalHeaderLabels(["Valeur actuelle", "Remplacement"])
        setup_table(self.replacements)
        self.replacements.setMaximumHeight(155)
        self.replacements.setMinimumHeight(100)
        form.addRow("Remplacements exacts", self.replacements)
        replacement_tools = QHBoxLayout()
        replacement_tools.addWidget(button("Ajouter", self.add_replacement))
        replacement_tools.addWidget(button("Retirer", self.remove_replacement))
        replacement_tools.addStretch()
        form.addRow(replacement_tools)
        self.replacement_tools = replacement_tools
        layout.addWidget(self.general)

        self.phone_group = QGroupBox("Téléphones")
        phone_form = QFormLayout(self.phone_group)
        from phonenumbers import SUPPORTED_REGIONS

        countries = {
            code: f"{QLocale('fr_' + code).nativeTerritoryName()} ({code})" for code in SUPPORTED_REGIONS
        }
        self.region = combo(dict(sorted(countries.items(), key=lambda pair: pair[1])), "CM")
        self.region.setToolTip("Pays utilisé pour reconnaître les numéros sans préfixe international.")
        self.phone_policy = combo(
            {
                "report": "Conserver et signaler",
                "first": "Garder le premier",
                "first_valid": "Garder le premier reconnu valide",
                "all": "Garder tous",
                "clear": "Vider la cellule",
            }
        )
        self.separator = QLineEdit("/")
        phone_form.addRow("Pays des numéros locaux", self.region)
        phone_form.addRow("Si plusieurs numéros sont présents", self.phone_policy)
        phone_form.addRow("Séparateur si tous conservés", self.separator)
        self.phone_policy.currentIndexChanged.connect(
            lambda: phone_form.setRowVisible(self.separator, self.phone_policy.currentData() == "all")
        )
        layout.addWidget(self.phone_group)

        self.number_group = QGroupBox("Nombres")
        numbers = QFormLayout(self.number_group)
        self.decimal_sep = combo({",": "Virgule", ".": "Point"})
        self.thousands_sep = combo({" ": "Espace", ",": "Virgule", ".": "Point", "": "Aucun"})
        self.decimal_sep.setToolTip("Exemple français : 1234,50 utilise une virgule décimale.")
        self.thousands_sep.setToolTip(
            "Séparateur présent dans la source : 1 234,50 utilise un espace. Aucun groupement ne sera ajouté en sortie."
        )
        self.currencies = QLineEdit("FCFA, XAF")
        numbers.addRow("Séparateur décimal en entrée", self.decimal_sep)
        numbers.addRow("Groupement en entrée", self.thousands_sep)
        numbers.addRow("Libellés monétaires à retirer", self.currencies)
        layout.addWidget(self.number_group)
        self.date_group = QGroupBox("Dates")
        date_form = QFormLayout(self.date_group)
        self.date_order = combo(
            {
                "dmy": "Jour / Mois / Année",
                "mdy": "Mois / Jour / Année",
                "ask": "Signaler les interprétations ambiguës",
            }
        )
        date_form.addRow("Interprétation des dates texte", self.date_order)
        layout.addWidget(self.date_group)

        chronology_group = QGroupBox("Relations entre dates")
        self.chronology_group = chronology_group
        chrono_layout = QVBoxLayout(chronology_group)
        self.chronology = QTableWidget(0, 3)
        self.chronology.setHorizontalHeaderLabels(
            ["Active", "Champ antérieur ou égal", "Champ postérieur ou égal"]
        )
        setup_table(self.chronology)
        self.chronology.setMaximumHeight(160)
        self.chronology.setColumnWidth(1, 300)
        chrono_layout.addWidget(self.chronology)
        chrono_tools = QHBoxLayout()
        chrono_tools.addWidget(button("Ajouter une relation", self.add_chronology))
        chrono_tools.addWidget(button("Supprimer la sélection", self.remove_chronology))
        chrono_tools.addStretch()
        chrono_layout.addLayout(chrono_tools)
        layout.addWidget(chronology_group)

        self.seg_group = QGroupBox("Segmentation entreprise / BP / téléphone / localisation")
        self.seg_group.setCheckable(True)
        self.seg_group.setChecked(False)
        seg_form = QFormLayout(self.seg_group)
        self.seg_boxes = {name: QComboBox() for name in ("source", "name", "bp", "phone", "location")}
        for name, label in (
            ("source", "Colonne composite"),
            ("name", "Destination entreprise"),
            ("bp", "Destination BP"),
            ("phone", "Destination téléphone"),
            ("location", "Destination localisation"),
        ):
            seg_form.addRow(label, self.seg_boxes[name])
        self.seg_overwrite = QCheckBox("Remplacer les destinations déjà renseignées")
        self.seg_overwrite.setToolTip(
            "Option destructive : remplacer une destination déjà remplie. Désactivée par défaut."
        )
        seg_form.addRow(self.seg_overwrite)
        self.seg_group.toggled.connect(
            lambda checked: [seg_form.setRowVisible(row, checked) for row in range(seg_form.rowCount())]
        )
        for row in range(seg_form.rowCount()):
            seg_form.setRowVisible(row, False)
        layout.addWidget(self.seg_group)
        layout.addStretch()
        self.field.currentIndexChanged.connect(self.change_field)

    def add_replacement(self, before="", after=""):
        row = self.replacements.rowCount()
        self.replacements.insertRow(row)
        self.replacements.setItem(row, 0, item(before if isinstance(before, str) else ""))
        self.replacements.setItem(row, 1, item(after))

    def remove_replacement(self):
        for row in sorted({i.row() for i in self.replacements.selectedIndexes()}, reverse=True):
            self.replacements.removeRow(row)

    def save_field(self):
        if not self.current_code:
            return
        replacements = {}
        for row in range(self.replacements.rowCount()):
            before = self.replacements.item(row, 0).text()
            after = self.replacements.item(row, 1).text()
            if before in replacements:
                raise ValueError("Une valeur ne peut avoir qu'un seul remplacement.")
            if before:
                replacements[before] = after
        self.app.config.rules[self.current_code] = FieldRule(
            normalize_spaces=self.normalize.isChecked(),
            convert=self.convert.isChecked(),
            decimal_separator=self.decimal_sep.currentData(),
            thousands_separator=self.thousands_sep.currentData(),
            currency_tokens=[v.strip() for v in self.currencies.text().split(",") if v.strip()],
            date_order=self.date_order.currentData(),
            phone_region=self.region.currentData(),
            phone_policy=self.phone_policy.currentData(),
            phone_separator=self.separator.text(),
            length_policy=self.length.currentData(),
            invalid_policy=self.invalid.currentData(),
            replacements=replacements,
        )

    def change_field(self):
        previous = self.current_code
        try:
            self.save_field()
        except Exception as exc:
            with QSignalBlocker(self.field):
                self.field.setCurrentIndex(self.field.findData(previous))
            self.app.show_error(exc)
            return
        self.current_code = self.field.currentData()
        if not self.current_code:
            return
        rule = self.app.config.rules.get(self.current_code, FieldRule())
        self.normalize.setChecked(rule.normalize_spaces)
        self.convert.setChecked(rule.convert)
        self.length.setCurrentIndex(self.length.findData(rule.length_policy))
        self.invalid.setCurrentIndex(self.invalid.findData(rule.invalid_policy))
        self.replacements.setRowCount(0)
        for before, after in rule.replacements.items():
            self.add_replacement(before, after)
        self.region.setCurrentIndex(self.region.findData(rule.phone_region))
        self.phone_policy.setCurrentIndex(self.phone_policy.findData(rule.phone_policy))
        self.separator.setText(rule.phone_separator)
        self.decimal_sep.setCurrentIndex(self.decimal_sep.findData(rule.decimal_separator))
        self.thousands_sep.setCurrentIndex(self.thousands_sep.findData(rule.thousands_separator))
        self.currencies.setText(", ".join(rule.currency_tokens))
        self.date_order.setCurrentIndex(self.date_order.findData(rule.date_order))
        spec = next((f for f in self.app.config.fields if f.code == self.current_code), None)
        self.general.layout().setRowVisible(self.length, bool(spec and spec.max_length))
        self.general.layout().setRowVisible(self.invalid, bool(spec and spec.pattern))
        self.phone_group.setVisible(bool(spec and spec.role == Role.PHONE))
        self.number_group.setVisible(bool(spec and spec.dtype in {DataType.INTEGER, DataType.DECIMAL}))
        self.date_group.setVisible(bool(spec and spec.dtype == DataType.DATE))
        self.phone_group.layout().setRowVisible(self.separator, rule.phone_policy == "all")

    def load(self):
        self.current_code = None
        self.general.layout().setRowVisible(self.replacements, self.app.advanced)
        self.general.layout().setRowVisible(self.replacement_tools, self.app.advanced)
        with QSignalBlocker(self.field):
            self.field.clear()
            mapped = {m.field_code for m in self.app.config.mappings if m.field_code}
            for spec in self.app.config.fields:
                if spec.code in mapped:
                    self.field.addItem(f"{spec.code} · {spec.label}", spec.code)
        self.change_field()
        self.chronology.setRowCount(0)
        for rule in self.app.config.chronology:
            self.add_chronology(rule)
        seg = self.app.config.segmentation
        headers = (self.app.sheet_info() or {}).get("headers", [])
        for name, box in self.seg_boxes.items():
            box.clear()
            box.addItem("Non utilisée", None)
            for col, header in enumerate(headers, 1):
                box.addItem(f"{get_column_letter(col)} · {display(header)}", col)
            value = seg.source_column if name == "source" else getattr(seg, f"{name}_column")
            box.setCurrentIndex(max(0, box.findData(value)))
        self.seg_group.setChecked(seg.enabled)
        self.seg_overwrite.setChecked(seg.overwrite)

    def add_chronology(self, rule=None):
        if not isinstance(rule, ChronologyRule):
            dates = [f for f in self.app.config.fields if f.dtype == DataType.DATE]
            if len(dates) < 2:
                return
            rule = ChronologyRule(before=dates[0].code, after=dates[1].code)
        options = {
            f.code: f"{f.code} · {f.label}" for f in self.app.config.fields if f.dtype == DataType.DATE
        }
        row = self.chronology.rowCount()
        self.chronology.insertRow(row)
        active = QCheckBox()
        active.setChecked(rule.enabled)
        self.chronology.setCellWidget(row, 0, active)
        self.chronology.setCellWidget(row, 1, combo(options, rule.before))
        self.chronology.setCellWidget(row, 2, combo(options, rule.after))

    def remove_chronology(self):
        for row in sorted({i.row() for i in self.chronology.selectedIndexes()}, reverse=True):
            self.chronology.removeRow(row)

    def collect(self):
        self.save_field()
        self.app.config.chronology = [
            ChronologyRule(
                before=self.chronology.cellWidget(row, 1).currentData(),
                after=self.chronology.cellWidget(row, 2).currentData(),
                enabled=self.chronology.cellWidget(row, 0).isChecked(),
            )
            for row in range(self.chronology.rowCount())
        ]
        self.app.config.segmentation = SegmentationRule(
            enabled=self.seg_group.isChecked(),
            source_column=self.seg_boxes["source"].currentData() or 1,
            name_column=self.seg_boxes["name"].currentData(),
            bp_column=self.seg_boxes["bp"].currentData(),
            phone_column=self.seg_boxes["phone"].currentData(),
            location_column=self.seg_boxes["location"].currentData(),
            overwrite=self.seg_overwrite.isChecked(),
        )


class ExportPage(QWidget):
    def __init__(self, app):
        super().__init__()
        self.app = app
        self.published_path = None
        layout = QVBoxLayout(self)
        self.summary = QLabel("Aucune analyse disponible")
        self.summary.setWordWrap(True)
        self.summary.setProperty("heading", True)
        layout.addWidget(self.summary)
        group = QGroupBox("Classeur de sortie")
        form = QFormLayout(group)
        self.path = QLineEdit()
        self.path.setObjectName("exportPath")
        form.addRow("Fichier Excel", ImportPage.file_row(self.path, self.choose_path))
        self.path.textChanged.connect(self.destination_changed)
        self.harmonized = QCheckBox("Ajouter Contrôle, Nom_original et À vérifier")
        self.harmonized.setChecked(True)
        form.addRow(self.harmonized)
        self.collision = combo(
            {
                "suffix": "Conserver les rapports existants et suffixer les nouveaux",
                "replace": "Remplacer les feuilles de rapport existantes",
            }
        )
        form.addRow("Noms de rapports déjà présents", self.collision)
        horizontal = QCheckBox("Orientation horizontale des cellules")
        horizontal.setChecked(True)
        horizontal.setEnabled(False)
        thousands = QCheckBox("Nombres sans séparateur de milliers")
        thousands.setChecked(True)
        thousands.setEnabled(False)
        form.addRow(horizontal)
        form.addRow(thousands)
        self.allow_issues = QCheckBox("Exporter avec les anomalies restantes signalées")
        self.allow_issues.setObjectName("allowIssues")
        form.addRow(self.allow_issues)
        layout.addWidget(group)
        self.sheets = QLabel()
        self.sheets.setWordWrap(True)
        layout.addWidget(self.sheets)
        self.harmonized.toggled.connect(self.show_sheets)
        self.collision.currentIndexChanged.connect(self.show_sheets)
        self.harmonized.toggled.connect(self.destination_changed)
        self.collision.currentIndexChanged.connect(self.destination_changed)
        self.export_button = button("Générer le classeur Excel", app.run_export, "exportButton")
        self.export_button.setProperty("primary", True)
        layout.addWidget(self.export_button)
        self.output = QLabel("")
        self.output.setWordWrap(True)
        self.output.setTextInteractionFlags(Qt.TextInteractionFlag.TextSelectableByMouse)
        layout.addWidget(self.output)
        opens = QHBoxLayout()
        self.open_file = button("Ouvrir le fichier", self.open_result)
        self.open_file.setIcon(icon("file-spreadsheet"))
        self.open_folder = button("Ouvrir le dossier", self.open_directory)
        self.open_folder.setIcon(icon("folder-open"))
        opens.addWidget(self.open_file)
        opens.addWidget(self.open_folder)
        opens.addStretch()
        layout.addLayout(opens)
        self.show_open_buttons()
        layout.addStretch()

    def destination_changed(self):
        self.published_path = None
        if hasattr(self, "output"):
            self.output.clear()
            self.show_open_buttons()
        if hasattr(self.app, "export_page"):
            self.app.update_navigation()

    def show_open_buttons(self):
        if hasattr(self, "open_file"):
            self.open_file.setVisible(bool(self.published_path))
            self.open_folder.setVisible(bool(self.published_path))

    def open_result(self):
        if self.published_path:
            QDesktopServices.openUrl(QUrl.fromLocalFile(self.published_path))

    def open_directory(self):
        if self.published_path:
            QDesktopServices.openUrl(QUrl.fromLocalFile(str(Path(self.published_path).parent)))

    def show_sheets(self):
        names = [s["name"] for s in (self.app.metadata or {}).get("sheets", [])]
        if self.harmonized.isChecked():
            for report in ["Contrôle", "Nom_original", "À vérifier"]:
                if report in names and self.collision.currentData() == "replace":
                    continue
                name, i = report, 2
                while name in names:
                    name = f"{report} ({i})"
                    i += 1
                names.append(name)
        self.sheets.setText("Feuilles du fichier final : " + " · ".join(names))

    def choose_path(self):
        path, _ = QFileDialog.getSaveFileName(
            self, "Enregistrer le résultat", self.path.text(), "Classeur Excel (*.xlsx)"
        )
        if path:
            self.path.setText(path if path.lower().endswith(".xlsx") else path + ".xlsx")

    def load(self):
        result = self.app.result
        self.summary.setText(
            f"{result.row_count} lignes analysées · {len(result.changes)} cellules modifiées · "
            f"{len(result.findings)} anomalies restantes"
            if result
            else "L'analyse doit être exécutée avant l'export."
        )
        if not self.path.text() and self.app.config.source_path:
            source = Path(self.app.config.original_path or self.app.config.source_path)
            self.path.setText(str(source.with_name(source.stem + " - traite.xlsx")))
        self.harmonized.setChecked(self.app.config.harmonized)
        self.collision.setCurrentIndex(self.collision.findData(self.app.config.report_collision))
        self.allow_issues.setVisible(bool(result and result.findings))
        self.export_button.setEnabled(result is not None)
        self.show_sheets()
        self.show_open_buttons()

    def collect(self):
        self.app.config.harmonized = self.harmonized.isChecked()
        self.app.config.report_collision = self.collision.currentData()
