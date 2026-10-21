from __future__ import annotations

import hashlib
import io
import posixpath
import re
import warnings
from pathlib import Path
from zipfile import BadZipFile, ZipFile

from defusedxml import ElementTree as ET
from openpyxl import load_workbook

from .models import Cancelled


class WorkbookError(ValueError):
    pass


def load_preserving(path, **options):
    """Restore raw numeric values which Excel styles mislabel as impossible date serials."""
    with warnings.catch_warnings(record=True) as caught:
        warnings.simplefilter("always", UserWarning)
        workbook = load_workbook(path, **options)
    broken_serials = False
    for warning in caught:
        if "outside the limits for dates" in str(warning.message):
            broken_serials = True
        else:
            warnings.warn(str(warning.message), warning.category, stacklevel=2)
    if broken_serials:
        ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
        rel_ns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
        with ZipFile(path) as archive:
            relationships = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
            targets = {r.attrib["Id"]: r.attrib["Target"] for r in relationships}
            structure = ET.fromstring(archive.read("xl/workbook.xml"))
            for entry in structure.findall("s:sheets/s:sheet", ns):
                target = targets[entry.attrib[f"{{{rel_ns}}}id"]]
                target = (
                    target.lstrip("/")
                    if target.startswith("/")
                    else posixpath.normpath(posixpath.join("xl", target))
                )
                sheet = workbook[entry.attrib["name"]]
                for _, node in ET.iterparse(io.BytesIO(archive.read(target)), events=("end",)):
                    if node.tag == f"{{{ns['s']}}}c":
                        coordinate = node.attrib.get("r")
                        numeric = node.find("s:v", ns)
                        if (
                            coordinate
                            and numeric is not None
                            and numeric.text
                            and node.attrib.get("t", "n") == "n"
                        ):
                            cell = sheet[coordinate]
                            if cell.data_type == "e" and cell.value == "#VALUE!":
                                cell.value = (
                                    int(numeric.text)
                                    if re.fullmatch(r"[+-]?\d+", numeric.text)
                                    else float(numeric.text)
                                )
                                cell.data_type = "n"
                        node.clear()
    return workbook


def sha256(path: str | Path) -> str:
    digest = hashlib.sha256()
    with Path(path).open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def preflight(path: str | Path) -> None:
    path = Path(path)
    if path.suffix.lower() != ".xlsx":
        raise WorkbookError("La première version accepte uniquement les fichiers .xlsx sans macros.")
    unsupported = (
        "xl/drawings/",
        "xl/pivotTables/",
        "xl/pivotCache/",
        "xl/slicers/",
        "xl/activeX/",
        "xl/embeddings/",
        "xl/ctrlProps/",
        "_xmlsignatures/",
    )
    try:
        with ZipFile(path) as archive:
            names = archive.namelist()
            if len(names) > 20000 or sum(i.file_size for i in archive.infolist()) > 1_000_000_000:
                raise WorkbookError(
                    "Classeur trop volumineux pour cette version (limite décompressée : 1 Go)."
                )
            if "xl/workbook.xml" not in names:
                raise WorkbookError("Le fichier n'est pas un classeur Excel valide.")
            blocked = [name for name in names if name.startswith(unsupported) or "vbaProject" in name]
            if blocked:
                raise WorkbookError("Objets non conservables dans cette version : " + ", ".join(blocked[:3]))
            for name in names:
                if name.startswith("xl/worksheets/") and name.endswith(".xml"):
                    content = archive.read(name)
                    if b"<!DOCTYPE" in content or b"<!ENTITY" in content:
                        raise WorkbookError("Déclarations XML DTD ou entités interdites dans le classeur.")
                    if re.search(rb"<(?:[\w.-]+:)?extLst(?:\s|/?>)", content):
                        raise WorkbookError(
                            "Extensions Excel avancées détectées. Export annulé pour les préserver."
                        )
    except BadZipFile as exc:
        raise WorkbookError("Fichier invalide, endommagé ou protégé par chiffrement.") from exc


def inspect_workbook(path: str | Path, header_row: int = 1, cancel=None) -> dict:
    check_cancel(cancel)
    preflight(path)
    workbook = load_preserving(path, data_only=False)
    sheets = []
    try:
        check_cancel(cancel)
        for sheet in workbook.worksheets:
            check_cancel(cancel)
            preview = list(
                sheet.iter_rows(
                    min_row=header_row, max_row=min(sheet.max_row, header_row + 8), values_only=True
                )
            )
            sheets.append(
                {
                    "name": sheet.title,
                    "rows": sheet.max_row,
                    "columns": sheet.max_column,
                    "headers": list(preview[0]) if preview else [],
                    "preview": preview[1:],
                }
            )
    finally:
        workbook.close()
    check_cancel(cancel)
    return {"path": str(path), "sheets": sheets}


def check_cancel(event) -> None:
    if event is not None and event.is_set():
        raise Cancelled("Opération annulée ; aucun fichier final n'a été écrit.")
