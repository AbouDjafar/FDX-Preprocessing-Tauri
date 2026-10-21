from __future__ import annotations

import os
import re
import tempfile
from collections import Counter
from copy import copy
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import AnalysisResult, DataType, ProjectConfig, display, equal_values
from .workbooks import check_cancel, load_preserving, sha256

REPORTS = ("Contrôle", "Nom_original", "À vérifier")


def remove_grouping(format_code: str) -> str:
    # Preserve literal currency/text, conditions, percent and date semantics.
    tokens = re.findall(r'"[^"]*"|\\.|\[[^\]]*\]|_.|\*.|.', format_code)
    for i, token in enumerate(tokens):
        if token in {",", " ", "\u00a0", "\u202f"}:
            left = tokens[i - 1] if i else ""
            right = tokens[i + 1] if i + 1 < len(tokens) else ""
            if left in {"#", "0", "?"} and right in {"#", "0", "?"}:
                tokens[i] = ""
    return "".join(tokens)


def _date_format(mask: str) -> str:
    return (
        mask.replace("AAAA", "yyyy")
        .replace("YYYY", "yyyy")
        .replace("JJ", "dd")
        .replace("DD", "dd")
        .replace("MM", "mm")
    )


def _write_value(cell, value) -> None:
    if isinstance(value, Decimal):
        value = int(value) if value == value.to_integral_value() else float(value)
    cell.value = value
    if isinstance(value, str):
        cell.data_type = "s"


def _report(workbook, name, headers, rows, collision):
    if name in workbook.sheetnames:
        if collision == "replace":
            del workbook[name]
        else:
            idx = 2
            while f"{name} ({idx})" in workbook.sheetnames:
                idx += 1
            name = f"{name} ({idx})"
    sheet = workbook.create_sheet(name)
    sheet.append(headers)
    for row_number, values in enumerate(rows, 2):
        sheet.append([display(v) if isinstance(v, (dict, list)) else v for v in values])
        for column in range(1, len(values) + 1):
            cell = sheet.cell(row_number, column)
            if isinstance(cell.value, str):
                cell.data_type = "s"
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="153E4A")
        cell.font = Font(name="Aptos", bold=True, color="FFFFFF")
        cell.alignment = Alignment(wrap_text=True, vertical="center", textRotation=0)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for i, header in enumerate(headers, 1):
        sheet.column_dimensions[get_column_letter(i)].width = min(65, max(14, len(header) + 3))
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    return sheet.title


def numeric_format(value, previous):
    tokens = re.findall(r'"[^\"]*"|\\.|\[[^\]]*\]|_.|\*.|.', previous)
    percent = "%" in tokens
    number = Decimal(str(value)) * (100 if percent else 1)
    places = max(0, -number.normalize().as_tuple().exponent)
    if places > 30 or (number and abs(number) >= Decimal("1e16")):
        return "0.##############E+00" + ("%" if percent else "")
    return "0" + ("." + "#" * places if places else "") + ("%" if percent else "")


def export_excel(
    config: ProjectConfig,
    result: AnalysisResult,
    destination: str | Path,
    allow_issues: bool = False,
    overwrite: bool = False,
    cancel=None,
    progress=None,
) -> Path:
    destination = Path(destination).expanduser().resolve()
    protected_paths = {Path(config.source_path).resolve()}
    if config.original_path:
        protected_paths.add(Path(config.original_path).resolve())
    if destination in protected_paths:
        raise ValueError("Choisissez un autre nom : le fichier source reste intact.")
    if destination.suffix.lower() != ".xlsx":
        raise ValueError("Le format de sortie doit être .xlsx.")
    if destination.exists() and not overwrite:
        raise FileExistsError(
            "Ce fichier existe déjà. Choisissez un autre nom ou confirmez son remplacement."
        )
    if not destination.parent.is_dir():
        raise ValueError("Le dossier de destination n'existe pas.")
    if config.fingerprint() != result.config_hash or sha256(config.source_path) != result.source_hash:
        raise ValueError("L'aperçu est obsolète : relancez l'analyse avant d'exporter.")
    if result.findings and not allow_issues:
        raise ValueError("Des anomalies restent présentes. Validez explicitement leur export.")
    check_cancel(cancel)
    if progress:
        progress(0, 4, "Préparation de la copie")
    wb = load_preserving(config.source_path, data_only=False, rich_text=True)
    temp_path = None
    source_sheet_names = list(wb.sheetnames)
    try:
        ws = wb[config.sheet]
        mapping = {
            m.column: next(f for f in config.fields if f.code == m.field_code)
            for m in config.mappings
            if m.field_code
        }
        for change in result.changes:
            _write_value(ws.cell(change.row, change.column), change.after)
        for col, field in mapping.items():
            for row in range(config.header_row + 1, ws.max_row + 1):
                cell = ws.cell(row, col)
                if cell.data_type == "f":
                    continue
                if isinstance(cell.value, (date, datetime)) and field.dtype == DataType.DATE:
                    cell.number_format = _date_format(field.date_format)
                elif isinstance(cell.value, str) and field.dtype == DataType.TEXT:
                    cell.number_format = "@"
        if config.harmonized:
            if config.report_collision == "replace" and config.sheet in REPORTS:
                raise ValueError("La feuille de données porte un nom réservé aux rapports.")
            issues_by_code = Counter(f.field_code for f in result.findings)
            changed_by_code = Counter(c.field_code for c in result.changes)
            blanked_by_code = Counter(
                c.field_code for c in result.changes if c.after is None or c.after == ""
            )
            columns = {m.field_code: m.column for m in config.mappings if m.field_code}
            _report(
                wb,
                "Contrôle",
                [
                    "Code",
                    "Colonne",
                    "Libellé",
                    "Stockage",
                    "Rôle",
                    "Taille max",
                    "Lignes analysées",
                    "Renseignées",
                    "Vides",
                    "Cellules non conformes avant",
                    "Cellules modifiées",
                    "Cellules vidées",
                    "Anomalies restantes",
                    "Longueur max",
                ],
                [
                    [
                        f.code,
                        columns[f.code],
                        f.label,
                        f.dtype.value,
                        f.role.value,
                        f.max_length,
                        result.row_count,
                        result.filled.get(f.code, 0),
                        result.row_count - result.filled.get(f.code, 0),
                        result.counts_before.get(f.code, 0),
                        changed_by_code[f.code],
                        blanked_by_code[f.code],
                        issues_by_code[f.code],
                        result.maximum_lengths.get(f.code, 0),
                    ]
                    for f in config.fields
                    if f.code in columns
                ],
                config.report_collision,
            )
            originals = [[o["sheet"], o["row"], o["column"], display(o["value"])] for o in result.originals]
            original_coordinates = {(o["row"], o["column"]) for o in result.originals}
            originals.extend(
                [config.sheet, c.row, c.column, display(c.before)]
                for c in result.changes
                if (c.row, c.column) not in original_coordinates
            )
            _report(
                wb,
                "Nom_original",
                ["Feuille", "Ligne", "Colonne", "Valeur source"],
                originals,
                config.report_collision,
            )
            _report(
                wb,
                "À vérifier",
                [
                    "Feuille",
                    "Ligne",
                    "Colonne",
                    "Code",
                    "Valeur source",
                    "Valeur actuelle",
                    "Catégorie",
                    "Problème",
                    "Proposition",
                ],
                [
                    [
                        config.sheet,
                        f.row,
                        f.column,
                        f.field_code,
                        display(f.original),
                        display(f.current),
                        f.kind,
                        f.message,
                        display(f.proposed),
                    ]
                    for f in result.findings
                ],
                config.report_collision,
            )
        check_cancel(cancel)
        if progress:
            progress(1, 4, "Formats horizontaux sans groupement des milliers")
        for sheet in wb.worksheets:
            for dimension in [*sheet.row_dimensions.values(), *sheet.column_dimensions.values()]:
                alignment = copy(dimension.alignment)
                alignment.textRotation = 0
                dimension.alignment = alignment
                dimension.number_format = remove_grouping(dimension.number_format)
            for conditional in sheet.conditional_formatting:
                for rule in sheet.conditional_formatting[conditional]:
                    if rule.dxf and rule.dxf.numFmt:
                        rule.dxf.numFmt.formatCode = remove_grouping(rule.dxf.numFmt.formatCode)
            for i, cell in enumerate(list(sheet._cells.values())):
                if i % 10000 == 0:
                    check_cancel(cancel)
                if cell.alignment.textRotation:
                    alignment = copy(cell.alignment)
                    alignment.textRotation = 0
                    cell.alignment = alignment
                if cell.data_type == "f":
                    cell.number_format = remove_grouping(cell.number_format)
                elif isinstance(cell.value, (int, float, Decimal)) and not isinstance(cell.value, bool):
                    # Only the display changes. Preserve the full stored fractional value.
                    cell.number_format = numeric_format(cell.value, cell.number_format)
        wb.calculation.fullCalcOnLoad = True
        wb.calculation.forceFullCalc = True
        with tempfile.NamedTemporaryFile(
            prefix=".fdx-", suffix=".xlsx", dir=destination.parent, delete=False
        ) as stream:
            temp_path = Path(stream.name)
        if progress:
            progress(2, 4, "Écriture du classeur temporaire")
        wb.save(temp_path)
        wb.close()
        check_cancel(cancel)
        if progress:
            progress(3, 4, "Réouverture et vérification de la conservation")
        _verify(config, result, temp_path, source_sheet_names, cancel)
        check_cancel(cancel)
        if sha256(config.source_path) != result.source_hash:
            raise ValueError("La source a changé pendant l'export. Le fichier final n'a pas été publié.")
        if destination.exists() and not overwrite:
            raise FileExistsError("Un fichier a été créé à cet emplacement pendant l'export.")
        os.replace(temp_path, destination)
        temp_path = None
        if progress:
            progress(4, 4, "Export vérifié")
        return destination
    finally:
        wb.close()
        if temp_path is not None:
            temp_path.unlink(missing_ok=True)


def _verify(config, result, path, original_names, cancel):
    output = load_workbook(path, data_only=False, rich_text=True)
    original = load_preserving(config.source_path, data_only=False, rich_text=True)
    try:
        changes = {(c.row, c.column): c.after for c in result.changes}
        for name in original_names:
            if config.harmonized and config.report_collision == "replace" and name in REPORTS:
                continue
            if name not in output.sheetnames:
                raise ValueError(f"Feuille perdue à l'export : {name}")
            before, after = original[name], output[name]
            if before.max_row != after.max_row or before.max_column != after.max_column:
                raise ValueError(f"Dimensions modifiées sans autorisation : {name}")
            if set(map(str, before.merged_cells.ranges)) != set(map(str, after.merged_cells.ranges)):
                raise ValueError(f"Fusions de cellules modifiées : {name}")
            if before.auto_filter.ref != after.auto_filter.ref or before.freeze_panes != after.freeze_panes:
                raise ValueError(f"Filtres ou volets modifiés : {name}")
            for i, ((row, col), cell) in enumerate(before._cells.items()):
                if i % 10000 == 0:
                    check_cancel(cancel)
                expected = changes.get((row, col), cell.value) if name == config.sheet else cell.value
                actual = after.cell(row, col).value
                if isinstance(expected, Decimal) and isinstance(actual, (float, int)):
                    matches = Decimal(str(actual)) == expected
                else:
                    matches = equal_values(expected, actual)
                if not matches:
                    raise ValueError(f"Valeur non conservée : {name}!{cell.coordinate}")
        for sheet in output.worksheets:
            for cell in sheet._cells.values():
                if cell.alignment.textRotation:
                    raise ValueError(f"Orientation verticale restante : {sheet.title}!{cell.coordinate}")
                if isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
                    if remove_grouping(cell.number_format) != cell.number_format:
                        raise ValueError("Séparateur de milliers restant à l'export.")
        if config.harmonized:
            name = (
                "À vérifier"
                if config.report_collision == "replace" or "À vérifier" not in original_names
                else None
            )
            if name and output[name].max_row - 1 != len(result.findings):
                raise ValueError("Le nombre d'anomalies exportées est incohérent.")
    finally:
        output.close()
        original.close()
