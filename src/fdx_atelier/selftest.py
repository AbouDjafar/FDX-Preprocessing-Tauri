"""Offline verification of both the installed runtime and the native interface."""

from __future__ import annotations

import json
import os
import tempfile
from pathlib import Path


def run(output: str, scale: float = 1.0):
    os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
    os.environ["QT_SCALE_FACTOR"] = str(scale)
    from openpyxl import Workbook, load_workbook
    from PySide6.QtWidgets import QApplication

    from .engine import analyze
    from .exporter import export_excel
    from .models import ColumnMapping, DataType, Decision, FieldRule, FieldSpec, ProjectConfig, Role, cell_key
    from .ui.app import MainWindow
    from .ui.session import decision_stamp
    from .ui.theme import apply_theme
    from .workbooks import inspect_workbook, sha256

    output = Path(output).resolve()
    output.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="fdx-selftest-") as temp:
        source = Path(temp) / "source.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Données"
        sheet.append(["Entreprise", "Téléphone", "Montant", "Date"])
        sheet.append(["Atelier Cameroun", "699123456/677123456", "1 234,50", "31/09/2025"])
        sheet.append(["Entreprise B", "+390498976845", 1250.125, "18 Fév. 2025"])
        workbook.save(source)
        workbook.close()
        fields = [
            FieldSpec(code="NOM", label="Entreprise"),
            FieldSpec(code="TEL", label="Téléphone", role=Role.PHONE),
            FieldSpec(code="MON", label="Montant", dtype=DataType.DECIMAL),
            FieldSpec(code="DATE", label="Date", dtype=DataType.DATE),
        ]
        cfg = ProjectConfig(
            source_path=str(source),
            original_path=str(source),
            sheet="Données",
            fields=fields,
            mappings=[
                ColumnMapping(column=i, field_code=f.code, confirmed=True) for i, f in enumerate(fields, 1)
            ],
            rules={f.code: FieldRule() for f in fields},
        )
        result = analyze(cfg)
        assert any(f.kind == "phone_multiple" for f in result.findings)
        assert any(f.kind == "date_invalid" for f in result.findings)
        fixed = analyze(
            cfg,
            {
                cell_key(cfg.sheet, 2, 2): Decision(action="set", value="699123456"),
                cell_key(cfg.sheet, 2, 4): Decision(action="set", value="2025-09-30"),
            },
        )
        assert not fixed.findings, [f.message for f in fixed.findings]
        target = Path(temp) / "résultat.xlsx"
        digest = sha256(source)
        export_excel(cfg, fixed, target)
        assert sha256(source) == digest
        wb = load_workbook(target)
        assert wb.active["C2"].value == 1234.5
        assert wb.active["B2"].value == "699123456"
        wb.close()
        app = QApplication.instance() or QApplication([])
        apply_theme(app)
        window = MainWindow()
        errors = []
        window.show_error = lambda error: errors.append(str(error))
        window.config = cfg
        window.metadata = inspect_workbook(source)
        window.import_page.load()
        window.show()
        screenshots = []
        viewports = (
            [(1920, 1080), (1366, 768), (1024, 768)]
            if scale == 1
            else [(int(1920 / scale), int(1080 / scale))]
        )
        for width, height in viewports:
            window.resize(width, height)
            for step in range(4):
                if step == 2:
                    result = analyze(window.config)
                    window.session.accept(result, decision_stamp(window.decisions))
                window.navigate(step)
                app.processEvents()
                name = output.parent / f"interface-{width}-{step + 1}.png"
                assert window.grab().save(str(name))
                screenshots.append(str(name))
            window.navigate(1)
            window.advanced_toggle.setChecked(True)
            window.prepare_page.tabs.setCurrentIndex(2)
            app.processEvents()
            window.grab().save(str(output.parent / f"avance-{width}.png"))
            window.advanced_toggle.setChecked(False)
        assert not errors, errors
        window.dirty = False
        window.close()
        app.processEvents()
        report = {
            "ok": True,
            "rows": result.row_count,
            "gui_steps": 4,
            "viewports": len(viewports),
            "scale": scale,
            "device_pixel_ratio": window.devicePixelRatioF(),
            "screenshots": screenshots,
            "engine": True,
            "pandera": True,
            "source_unchanged": True,
            "export_verified": True,
        }
        output.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
        return report
